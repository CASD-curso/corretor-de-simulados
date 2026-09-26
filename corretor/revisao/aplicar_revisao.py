"""
Le a planilha de revisao (corretor.revisao.gerar_planilha_revisao) depois
que o operador preencheu, e aplica as correcoes por cima dos dados
originais -- antes do calculo de nota em corretor.relatorio.gerar_excel.

O calculo estatistico nunca deve ler a saida do modelo direto: sempre passa
por aqui primeiro quando houve revisao (item 5 do plano).

Folha marcada com SIM em "Excluir?" no quadro de inscricoes repetidas do
Checkup (item 12) sai do resultado, e os itens dela na aba Correcao sao
ignorados.
"""
import openpyxl

from corretor.config import NUM_DIGITOS_INSCRICAO, NUM_QUESTOES
from corretor.revisao.gerar_planilha_revisao import CAMPO_INSCRICAO_REPETIDA, VALOR_EXCLUIR


def _normalizar_resposta_questao(valor):
    """
    Só normaliza caixa e espaços. Antes desta função também convertia o
    código curto do operador ('BRANCO'/'NULA') para o código longo do
    modelo ('EM BRANCO'/'NULA(MARCADAS>1)') -- essa conversão não existe
    mais porque agora os dois lados têm nomes distintos por construção:
    o modelo grava 'BRANCOM'/'MULTM' (sufixo M = decisão automática, sem
    revisão) e o operador digita 'BRANCO'/'MULT' (sem M = confirmado por
    humano). Cada um já é o valor final por si só, sem precisar de mapa.
    """
    return valor.strip().upper()


def _normalizar_digito_inscricao(valor, simulado, campo):
    """
    Quando o operador digita só um dígito solto (0-9) numa célula do Excel
    sem formatação de texto, o Excel guarda isso como NÚMERO, não como
    texto. openpyxl(data_only=True) então devolve um float, e
    str(1.0).strip() é "1.0" -- 3 caracteres, não 1. Sem essa limpeza, essa
    string de 3 caracteres é escrita em UMA posição do número de
    inscrição, e o ".0" sobra ali dentro ao juntar tudo de novo, deslocando
    e corrompendo as posições vizinhas (ex.: "2601105" virava algo como
    "26011.005").
    """
    valor = valor.strip().upper()
    if valor.endswith(".0") and valor[:-2].isdigit():
        valor = valor[:-2]
    if len(valor) != 1:
        print(f"AVISO: '{simulado}' campo '{campo}' -- valor corrigido inesperado "
              f"'{valor}' (esperado 1 caractere). Usando '?' para não corromper "
              f"as posições vizinhas.")
        valor = "?"
    return valor


def _normalizar_inscricao_inteira(valor, simulado):
    """
    Inscricao digitada inteira no item de inscricao repetida. A celula e
    formatada como texto, mas se o Excel ainda assim guardar numero, volta
    como float ('2601105.0') e sem o zero da frente -- os dois sao
    desfeitos aqui. Devolve None se nao sobrar um numero valido.
    """
    valor = str(valor).strip()
    if valor.endswith(".0") and valor[:-2].isdigit():
        valor = valor[:-2]
    if valor.isdigit() and len(valor) < NUM_DIGITOS_INSCRICAO:
        valor = valor.zfill(NUM_DIGITOS_INSCRICAO)
    if len(valor) != NUM_DIGITOS_INSCRICAO or not valor.isdigit():
        print(f"AVISO: '{simulado}' -- inscrição corrigida '{valor}' inválida "
              f"(esperado {NUM_DIGITOS_INSCRICAO} dígitos). Mantida a leitura original.")
        return None
    return valor


def _ler_excluidos(ck):
    """
    Simulados marcados com SIM no quadro de inscricoes repetidas do Checkup.
    O quadro e achado pelo cabecalho "Inscrição repetida"; as colunas
    "Simulado N" trazem o nome da folha, e a coluna seguinte a cada uma e o
    "Excluir?" dela.
    """
    excluidos = set()
    linhas = list(ck.iter_rows(values_only=True))
    for i, row in enumerate(linhas):
        if "Inscrição repetida" not in row:
            continue
        col_numero = row.index("Inscrição repetida")
        cols_nome = [j for j, titulo in enumerate(row)
                     if isinstance(titulo, str) and titulo.startswith("Simulado ")]
        for dados in linhas[i + 1:]:
            if col_numero >= len(dados) or not dados[col_numero]:
                break
            for j in cols_nome:
                simulado, excluir = dados[j], dados[j + 1]
                if simulado and str(excluir or "").strip().upper() == VALOR_EXCLUIR:
                    excluidos.add(simulado)
        break
    return excluidos


def _recalcular_link_revisao(linha, simulados_com_pendencia_restante):
    """
    Depois da correcao, "BRANCOM" ou "MULTM" que o operador reescreveu como
    "BRANCO"/"MULT" sao respostas legitimas que ele CONFIRMOU ao revisar --
    o proprio texto final ja denuncia isso (perdeu o sufixo M). O que
    decide se a linha ainda precisa de atencao e se ela tem algum item que
    ficou sem revisao nenhuma (Corrigido em branco na aba Correcao, ou nem
    reescaneada nem transcrita no Checkup) -- nao o valor final em si.

    'Não' em branco (nunca chegou a passar pela extracao) e "?"/"NAO_LIDO"
    persistentes continuam sinalizando erro como rede de seguranca, pro
    caso de um item nem ter entrado na planilha de revisao por algum motivo.
    """
    inscricao = linha.get("Inscricao", "")
    alinhamento_nao_resolvido = linha.get("Alinhamento_OK", "Sim") == "Não"
    tem_erro = (
        linha["Simulado"] in simulados_com_pendencia_restante
        or alinhamento_nao_resolvido
        or "?" in inscricao
        or "NAO_LIDO" in inscricao
    )
    if tem_erro:
        url = f"https://drive.google.com/drive/search?q={linha['Simulado']}.tiff"
        linha["Link_Revisao"] = f'=HYPERLINK("{url}"; "Ver Cartão")'
    else:
        linha["Link_Revisao"] = ""
    return linha


def aplicar_revisao(caminho_planilha_revisada, linhas_originais):
    """
    linhas_originais: a lista de dicts que gerar_planilha_unificada()
    devolve (o primeiro item da tupla) -- a mesma que iria direto pra
    gerar_excel_final() se não houvesse revisão. Retorna uma NOVA lista já
    com as correções aplicadas; linhas_originais não é alterada.
    """
    wb = openpyxl.load_workbook(caminho_planilha_revisada, data_only=True)
    por_simulado = {linha["Simulado"]: dict(linha) for linha in linhas_originais}

    # -------- aba Checkup: folhas excluidas no quadro de repetidas --------
    # Lido antes da aba Correcao, porque os itens dessas folhas sao ignorados.
    excluidos = _ler_excluidos(wb["Checkup"]) if "Checkup" in wb.sheetnames else set()

    # -------- aba Correção: aplica por (Simulado, Campo) --------
    pendentes = 0
    simulados_com_pendencia_restante = set()
    if "Correção" in wb.sheetnames:
        cr = wb["Correção"]
        for row in cr.iter_rows(min_row=3, values_only=True):
            simulado, campo, _motivo, corrigido = row[:4]
            if not simulado or simulado in excluidos:
                continue
            if campo == CAMPO_INSCRICAO_REPETIDA:
                # Em branco aqui e resposta valida: a folha estava certa,
                # quem foi lida errado e a outra da repeticao.
                if corrigido is None or str(corrigido).strip() == "":
                    continue
                if simulado in por_simulado:
                    nova = _normalizar_inscricao_inteira(corrigido, simulado)
                    if nova:
                        por_simulado[simulado]["Inscricao"] = nova
                continue
            if corrigido is None or str(corrigido).strip() == "":
                pendentes += 1
                simulados_com_pendencia_restante.add(simulado)
                continue
            corrigido = str(corrigido).strip()
            if simulado not in por_simulado:
                print(f"AVISO: '{simulado}' está na planilha de revisão mas não nos dados "
                      f"originais — ignorado.")
                continue
            linha = por_simulado[simulado]
            if campo.startswith("Q"):
                linha[campo] = _normalizar_resposta_questao(corrigido)
            elif campo.startswith("pos"):
                pos = int(campo[3:])
                digito_corrigido = _normalizar_digito_inscricao(corrigido, simulado, campo)
                inscricao = list(linha.get("Inscricao", "").ljust(NUM_DIGITOS_INSCRICAO, "?"))
                inscricao[pos] = digito_corrigido
                linha["Inscricao"] = "".join(inscricao)

    # -------- aba Checkup: folhas descartadas no alinhamento --------
    nao_resolvidos = []
    if "Checkup" in wb.sheetnames:
        ck = wb["Checkup"]
        capturando = False
        for row in ck.iter_rows(values_only=True):
            if row[0] == "Simulado" and row[1] == "Link_Original":
                capturando = True
                continue
            if not capturando:
                continue
            simulado = row[0]
            if not simulado:
                break  # fim da seção de alinhamento
            _link, inscricao_manual, respostas_manual = row[1], row[2], row[3]
            if not inscricao_manual or not respostas_manual:
                nao_resolvidos.append(simulado)
                continue
            respostas_manual = str(respostas_manual).strip().upper()
            if len(respostas_manual) != NUM_QUESTOES:
                print(f"AVISO: '{simulado}' tem Respostas_Manual com "
                      f"{len(respostas_manual)} caractere(s), esperado {NUM_QUESTOES} — ignorado.")
                nao_resolvidos.append(simulado)
                continue
            nova_linha = {
                "Simulado": simulado,
                "Inscricao": str(inscricao_manual).strip(),
                "Alinhamento_OK": "Não (revisado manualmente)",
                "Motivo_Alinhamento": "",
                "Bolhas_Ilegiveis": "",
            }
            for i, letra in enumerate(respostas_manual):
                # "BRANCO" (sem M) porque essa linha inteira já é
                # transcrição manual do operador olhando o scan -- não
                # existe leitura automática do modelo aqui para ter o
                # sufixo M.
                nova_linha[f"Q{i + 1}"] = letra if letra in "ABCDE" else "BRANCO"
            por_simulado[simulado] = nova_linha

    if pendentes:
        print(f"AVISO: ainda restam {pendentes} item(ns) sem revisão na aba 'Correção'. "
              f"Serão calculados com a leitura original do modelo.")
    if nao_resolvidos:
        print(f"AVISO: {len(nao_resolvidos)} folha(s) descartada(s) no alinhamento ainda sem "
              f"solução (nem reescaneada, nem transcrita): {nao_resolvidos}")

    if excluidos:
        print(f"{len(excluidos)} simulado(s) excluído(s) no quadro de inscrições repetidas: "
              f"{sorted(excluidos)}")

    linhas_finais = [
        _recalcular_link_revisao(linha, simulados_com_pendencia_restante)
        for simulado, linha in por_simulado.items()
        if simulado not in excluidos
    ]
    return linhas_finais
